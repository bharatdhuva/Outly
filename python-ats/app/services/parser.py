from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "application/octet-stream",
}

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MULTI_SPACE = re.compile(r"[ \t]{3,}")
MULTI_NEWLINE = re.compile(r"\n{4,}")


class ParseError(ValueError):
    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    content: str
    extension: str


def sanitise_text(text: str, max_chars: int) -> str:
    cleaned = text.replace("\x00", " ")
    cleaned = CONTROL_CHARS.sub(" ", cleaned)
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = MULTI_SPACE.sub("  ", cleaned)
    cleaned = MULTI_NEWLINE.sub("\n\n\n", cleaned)
    cleaned = cleaned.strip()
    if len(cleaned) > max_chars:
        cleaned = cleaned[:max_chars]
    return cleaned


def _extension(filename: str) -> str:
    name = filename.lower()
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


def validate_upload(
    filename: str,
    content_type: str | None,
    size: int,
    max_bytes: int,
) -> str:
    if size > max_bytes:
        raise ParseError(
            f"File too large. Maximum allowed size is {max_bytes // (1024 * 1024)} MB."
        )
    ext = _extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ParseError(
            "Unsupported file type. Please upload a .txt, .pdf, or .docx file."
        )
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime and mime not in ALLOWED_MIME_TYPES:
        raise ParseError("File content does not match allowed types.")
    return ext


def parse_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("Failed to parse PDF document. The file may be corrupt.") from exc

    if getattr(reader, "is_encrypted", False):
        raise ParseError("Encrypted PDFs are not supported. Please upload an unlocked file.")

    parts: list[str] = []
    try:
        for page in reader.pages:
            parts.append(page.extract_text() or "")
    except Exception as exc:  # noqa: BLE001
        raise ParseError("Failed to parse PDF document. The file may be corrupt.") from exc
    return "\n".join(parts)


def parse_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("Failed to parse DOCX document. The file may be corrupt.") from exc
    paragraphs = [p.text for p in document.paragraphs]
    table_cells: list[str] = []
    for table in document.tables:
        for row in table.rows:
            table_cells.extend(cell.text for cell in row.cells)
    return "\n".join([*paragraphs, *table_cells])


def parse_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_bytes(
    data: bytes,
    filename: str,
    content_type: str | None,
    max_bytes: int,
    max_chars: int,
) -> ParsedDocument:
    if not data:
        raise ParseError("The uploaded file is empty.")
    ext = validate_upload(filename, content_type, len(data), max_bytes)

    if ext == ".pdf":
        raw = parse_pdf(data)
    elif ext == ".docx":
        raw = parse_docx(data)
    else:
        raw = parse_txt(data)

    content = sanitise_text(raw, max_chars)
    if not content:
        raise ParseError(
            "No extractable text was found. Scanned or image-only PDFs are not supported. "
            "Please upload a text-based resume."
        )
    return ParsedDocument(filename=filename, content=content, extension=ext)
